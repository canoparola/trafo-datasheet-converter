"""
Transformer Datasheet Converter
Converts project-specific CSV to generic Excel format
"""

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import re
from config import (
    INPUT_FILES, OUTPUT_FILE, NOTES_PATH, HOLDS_PATH,
    DATASHEET_SECTIONS, CSV_COLUMNS, EXCEL_FORMAT, CSV_ENCODING,
    CLEANUP
)

class TransformerDatasheetConverter:
    def __init__(self):
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)  # Remove default sheet
        self.project_data = {}
        self.datasheet_data = []
        
    def read_csv_files(self):
        """Read and parse CSV input files"""
        print("[1/5] Reading CSV files...")
        
        all_data = []
        for file_path in INPUT_FILES:
            if os.path.exists(file_path):
                print(f"  Reading: {file_path}")
                df = pd.read_csv(file_path, encoding=CSV_ENCODING)
                all_data.append(df)
            else:
                print(f"  WARNING: File not found: {file_path}")
        
        if all_data:
            self.datasheet_data = pd.concat(all_data, ignore_index=True)
            print(f"  Total rows read: {len(self.datasheet_data)}")
        else:
            print("  ERROR: No CSV files found!")
            
    def clean_datasheet_data(self):
        """Clean and process datasheet data"""
        print("[2/5] Cleaning and converting to generic format...")
        
        # Remove completely empty rows
        self.datasheet_data = self.datasheet_data.dropna(how='all')
        
        # Fill NaN values
        self.datasheet_data.fillna('', inplace=True)
        
        # Remove leading/trailing whitespace
        if CLEANUP['remove_trailing_whitespace']:
            for col in self.datasheet_data.columns:
                if self.datasheet_data[col].dtype == 'object':
                    self.datasheet_data[col] = self.datasheet_data[col].str.strip()
        
        # Convert to generic format (remove project-specific values, Keep only structure)
        self.datasheet_data = self._convert_to_generic(self.datasheet_data)
        
        print(f"  Rows after cleaning: {len(self.datasheet_data)}")
    
    def _convert_to_generic(self, df):
        """Convert project-specific datasheet to generic template"""
        
        # Get column indices
        cols = df.columns.tolist()
        
        # Expected columns: SL. NO, REV, TECHNICAL PARTICULARS, PURCHASER'S REQUIREMENTS, UNIT, VENDOR'S RESPONSE
        processed_rows = []
        
        for idx, row in df.iterrows():
            new_row = {}
            
            # Copy SL.NO and REV as-is
            for i, col in enumerate(cols[:2]):
                new_row[col] = row[col] if i < len(row) else ''
            
            # TECHNICAL PARTICULARS - keep the description
            if len(cols) > 2:
                tech_particular = str(row[cols[2]]) if len(row) > 2 else ''
                new_row[cols[2]] = tech_particular
            
            # PURCHASER'S REQUIREMENTS - CLEAR PROJECT-SPECIFIC VALUES
            # Keep only if it's a note/hold reference, otherwise empty
            if len(cols) > 3:
                purchaser_req = str(row[cols[3]]) if len(row) > 3 else ''
                
                # Keep only structural notes (Hold Note X, See Note X, VTA markers)
                if any(keyword in purchaser_req.upper() for keyword in ['HOLD NOTE', 'SEE NOTE', 'VTA', 'TBD', 'NOTE']):
                    new_row[cols[3]] = purchaser_req
                else:
                    # Remove project-specific values, keep empty for user to fill
                    new_row[cols[3]] = ''
            
            # UNIT - keep as-is
            if len(cols) > 4:
                unit = str(row[cols[4]]) if len(row) > 4 else ''
                new_row[cols[4]] = unit
            
            # VENDOR'S RESPONSE - always empty (template for vendor to fill)
            if len(cols) > 5:
                new_row[cols[5]] = ''
            
            processed_rows.append(new_row)
        
        return pd.DataFrame(processed_rows)
    
    def read_notes_and_holds(self):
        """Read project notes and hold notes"""
        print("[3/5] Reading Notes and Holds...")
        
        self.project_notes = ""
        self.hold_notes = ""
        
        if os.path.exists(NOTES_PATH):
            with open(NOTES_PATH, 'r', encoding=CSV_ENCODING) as f:
                self.project_notes = f.read()
            print(f"  Project Notes loaded ({len(self.project_notes)} chars)")
        
        if os.path.exists(HOLDS_PATH):
            with open(HOLDS_PATH, 'r', encoding=CSV_ENCODING) as f:
                self.hold_notes = f.read()
            print(f"  Hold Notes loaded ({len(self.hold_notes)} chars)")
    
    def create_sheet1_header(self):
        """Create Sheet 1: Header (Template - to be filled by user)"""
        print("[4a/5] Creating Sheet 1 (Header - Template)...")
        
        ws = self.workbook.create_sheet(title="Header", index=0)
        
        # Title styling
        title_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        title_font = Font(bold=True, size=14, color="FFFFFF")
        
        # Header template
        ws['A1'] = "TRANSFORMER DATASHEET - GENERIC TEMPLATE"
        ws['A1'].font = title_font
        ws['A1'].fill = title_fill
        ws.merge_cells('A1:F1')
        
        ws['A3'] = "Project Name:"
        ws['B3'] = "___________________"
        
        ws['A4'] = "Equipment Tag:"
        ws['B4'] = "___________________"
        
        ws['A5'] = "Date:"
        ws['B5'] = "___________________"
        
        ws['A7'] = "This is a GENERIC template for transformer datasheet preparation."
        ws['A8'] = "Fill in the requirements per your project needs."
        
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 35
        
    def create_sheet2_notes_and_scope(self):
        """Create Sheet 2: Notes, Reference Docs, Scope"""
        print("[4b/5] Creating Sheet 2 (Notes and Scope)...")
        
        ws = self.workbook.create_sheet(title="Notes & Scope")
        
        # Header styling
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, size=12, color="FFFFFF")
        
        # Title
        ws['A1'] = "PROJECT NOTES & SCOPE"
        ws['A1'].font = header_font
        ws['A1'].fill = header_fill
        ws.merge_cells('A1:D1')
        
        # Notes content
        row = 2
        for line in self.project_notes.split('\n'):
            if line.strip():
                ws[f'A{row}'] = line
                row += 1
            elif CLEANUP['keep_single_blank_lines']:
                row += 1
        
        ws.column_dimensions['A'].width = 80
        
    def create_sheet3_holds_and_legends(self):
        """Create Sheet 3: Holds and Legends"""
        print("[4c/5] Creating Sheet 3 (Holds and Legends)...")
        
        ws = self.workbook.create_sheet(title="Holds & Legends")
        
        # Header styling
        header_fill = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")
        header_font = Font(bold=True, size=12, color="FFFFFF")
        
        # Title
        ws['A1'] = "HOLD NOTES"
        ws['A1'].font = header_font
        ws['A1'].fill = header_fill
        ws.merge_cells('A1:D1')
        
        # Hold notes content
        row = 2
        for line in self.hold_notes.split('\n'):
            if line.strip():
                ws[f'A{row}'] = line
                row += 1
            elif CLEANUP['keep_single_blank_lines']:
                row += 1
        
        # Legends
        row += 2
        ws[f'A{row}'] = "LEGENDS & ABBREVIATIONS"
        ws[f'A{row}'].font = header_font
        ws[f'A{row}'].fill = header_fill
        ws.merge_cells(f'A{row}:D{row}')
        
        row += 1
        legends = [
            "VTA: Vendor To Answer",
            "TBD: To Be Determined",
            "kVA: Kilovolt-Amperes",
            "CT: Current Transformer",
            "DGA: Dissolved Gas Analysis",
            "IEC: International Electrotechnical Commission",
        ]
        
        for legend in legends:
            ws[f'A{row}'] = legend
            row += 1
        
        ws.column_dimensions['A'].width = 80
    
    def create_sheet4_generic_datasheet(self):
        """Create Sheet 4: Generic Datasheet Template"""
        print("[5/5] Creating Sheet 4 (Generic Datasheet Template)...")
        
        ws = self.workbook.create_sheet(title="Datasheet")
        
        # Define styles
        header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        header_font = Font(bold=True, size=11, color="FFFFFF")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Column headers
        headers = ['SL. NO', 'REV', 'TECHNICAL PARTICULARS', "PURCHASER'S REQUIREMENTS", 'UNIT', "VENDOR'S RESPONSE"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            cell.border = border
        
        # Set column widths
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 8
        ws.column_dimensions['C'].width = 35
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 30
        
        # Set row height for header
        ws.row_dimensions[1].height = 30
        
        # Write data
        current_row = 2
        for idx, row_data in self.datasheet_data.iterrows():
            # Set row height for better readability
            ws.row_dimensions[current_row].height = None  # Auto height
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col)
                cell.value = value
                cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
                cell.border = border
            
            current_row += 1
        
        # Freeze panes (header row)
        ws.freeze_panes = 'A2'
    
    def save_workbook(self):
        """Save workbook to Excel file"""
        print(f"Saving to {OUTPUT_FILE}...")
        
        # Create output directory if it doesn't exist
        output_dir = os.path.dirname(OUTPUT_FILE)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir)
        
        self.workbook.save(OUTPUT_FILE)
        print(f"✓ Excel file created: {OUTPUT_FILE}")
        
    def run(self):
        """Run the complete conversion process"""
        print("\n" + "="*60)
        print("Transformer Datasheet Converter - GENERIC TEMPLATE")
        print("="*60 + "\n")
        
        self.read_csv_files()
        self.clean_datasheet_data()
        self.read_notes_and_holds()
        self.create_sheet1_header()
        self.create_sheet2_notes_and_scope()
        self.create_sheet3_holds_and_legends()
        self.create_sheet4_generic_datasheet()
        self.save_workbook()
        
        print("\n" + "="*60)
        print("✓ Conversion completed successfully!")
        print("✓ Generic template ready for use")
        print("="*60 + "\n")


if __name__ == "__main__":
    converter = TransformerDatasheetConverter()
    converter.run()
